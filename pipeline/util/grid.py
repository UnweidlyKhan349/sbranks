"""Helpers for reading spreadsheet tabs as plain 2-D grids.

    from pipeline.util.grid import load_grids
    grids = load_grids(t.raw("results.xlsx"))      # {tab title: Grid}
    g = grids["RR Scoring"]
    for r, c in g.find(r"^Team$"): ...
    g.cell(r, c)          # value or None (0-based row/col)
    g.text(r, c)          # cleaned string ('' for blank)

Values are cached results of formulas (``data_only=True``). Merged ranges are filled with
the top-left value when ``fill_merged=True`` (useful for bracket layouts).
"""
from __future__ import annotations

import re
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

import openpyxl

from ..schema import clean_name, num


@dataclass
class Grid:
    title: str
    rows: list[list[Any]]
    hidden: bool = False

    @property
    def nrows(self) -> int:
        return len(self.rows)

    @property
    def ncols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def cell(self, r: int, c: int) -> Any:
        if 0 <= r < len(self.rows) and 0 <= c < len(self.rows[r]):
            return self.rows[r][c]
        return None

    def text(self, r: int, c: int) -> str:
        return clean_name(self.cell(r, c))

    def num(self, r: int, c: int) -> float | None:
        return num(self.cell(r, c))

    def row_texts(self, r: int) -> list[str]:
        return [clean_name(v) for v in (self.rows[r] if 0 <= r < len(self.rows) else [])]

    def find(self, pattern: str, flags: int = re.I) -> Iterator[tuple[int, int]]:
        """Yield (row, col) of cells whose cleaned text matches ``pattern`` (re.search)."""
        rx = re.compile(pattern, flags)
        for r, row in enumerate(self.rows):
            for c, v in enumerate(row):
                if v is not None and rx.search(clean_name(v)):
                    yield r, c

    def find_first(self, pattern: str, flags: int = re.I) -> tuple[int, int] | None:
        return next(self.find(pattern, flags), None)

    def dump(self, max_rows: int = 40, max_cols: int = 30, width: int = 14) -> str:
        """Human-readable preview (handy while writing a parser)."""
        lines = []
        for r in range(min(self.nrows, max_rows)):
            cells = [clean_name(self.cell(r, c))[:width].ljust(width) for c in range(min(self.ncols, max_cols))]
            lines.append(f"{r:3d}| " + "|".join(cells).rstrip())
        return "\n".join(lines)


def load_grids(path: Path | str, fill_merged: bool = False) -> dict[str, Grid]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True, read_only=not fill_merged)
    out: dict[str, Grid] = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # trim trailing empty rows
        while rows and all(v is None or str(v).strip() == "" for v in rows[-1]):
            rows.pop()
        if fill_merged and hasattr(ws, "merged_cells"):
            for rng in ws.merged_cells.ranges:
                v = ws.cell(rng.min_row, rng.min_col).value
                for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                    for c in range(rng.min_col - 1, rng.max_col):
                        while len(rows[r]) <= c:
                            rows[r].append(None)
                        if rows[r][c] is None:
                            rows[r][c] = v
        out[ws.title] = Grid(ws.title, rows, hidden=getattr(ws, "sheet_state", "visible") != "visible")
    wb.close()
    return out


def header_index(headers: list[str], *patterns: str) -> int | None:
    """Index of the first header matching any regex in ``patterns`` (case-insensitive)."""
    for pat in patterns:
        rx = re.compile(pat, re.I)
        for i, h in enumerate(headers):
            if h and rx.search(h):
                return i
    return None
